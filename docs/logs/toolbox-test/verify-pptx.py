# 用 python-pptx 读回生成的 pptx，逐页校验标题/要点/备注；再用 zipfile 校验 ZIP 结构
import sys, zipfile, io
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

path = sys.argv[1]
print("=== ZIP 结构校验 ===")
z = zipfile.ZipFile(path)
bad = z.testzip()
print("testzip()（坏条目）:", bad)
names = z.namelist()
print("条目数:", len(names))
print("[Content_Types].xml 在根:", "[Content_Types].xml" in names, "| 第一个条目:", names[0])
print("根目录下的条目:", [n for n in names if "/" not in n])
for n in ["_rels/.rels", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"]:
    print("  存在", n, ":", n in names)
# 每个 xml 部件能否被 xml 解析（结构完好性）
import xml.etree.ElementTree as ET
for n in names:
    if n.endswith(".xml") or n.endswith(".rels"):
        ET.fromstring(z.read(n))
print("所有 XML/rels 部件均可被 XML 解析: True")

print()
print("=== python-pptx 读回 ===")
prs = Presentation(path)
print("幻灯片尺寸:", prs.slide_width, "x", prs.slide_height, "(16:9 期望 12192000 x 6858000)")
print("页数:", len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    print(f"--- 第 {i} 页 ---")
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        ph = shape.placeholder_format
        kind = f"{ph.type}" if ph is not None else "non-ph"
        texts = [p.text for p in shape.text_frame.paragraphs]
        print(f"  [{kind}] " + " | ".join(texts))
    if slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame.text
        print("  (备注) " + repr(nt))

# 逐页断言
slides = list(prs.slides)
def texts_of(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append("\n".join(p.text for p in shape.text_frame.paragraphs))
    return out
assert len(slides) == 4, f"页数应为 4，实际 {len(slides)}"
assert slides[0].shapes.title.text == "Android 上的真 PPTX 生成"
assert "无 npm 依赖" in texts_of(slides[0])[1]
assert slides[1].shapes.title.text == "为什么要自己写 ZIP"
body = texts_of(slides[1])[1]
assert "装不了 npm 包" in body and "所以容器、CRC32、OOXML 全自己实现" in body
assert slides[2].shapes.title.text == "最小合法 OOXML 集合"
assert slides[3].shapes.title.text == "中文支持"
b3 = texts_of(slides[3])[1]
assert "一级要点：ea typeface = 等线" in b3 and "二级要点：缩进 685800 EMU" in b3
assert "带\x0b换行的要点\x0b也要正常" in b3, repr(b3)   # python-pptx 用 \v 表示段内 <a:br/>
import re as _re
_s4 = zipfile.ZipFile(path).read("ppt/slides/slide4.xml").decode("utf-8")
assert "<a:br/>" in _s4, "段内换行必须落成真正的 <a:br/> 元素"
assert slides[1].notes_slide.notes_text_frame.text == "备注：这一段讲环境约束。"
assert slides[3].notes_slide.notes_text_frame.text == "备注第二页\n第二行"
print()
print("全部断言通过 ✅")
