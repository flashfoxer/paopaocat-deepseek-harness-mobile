from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation()
prs.slide_width = 12192000
prs.slide_height = 6858000
layout = prs.slide_layouts[1]  # Title and Content
s = prs.slides.add_slide(layout)
s.shapes.title.text = "小节标题"
body = s.placeholders[1].text_frame
body.text = "要点1"
p = body.add_paragraph(); p.text = "要点2"
s.notes_slide.notes_text_frame.text = "演讲者备注"
prs.save(r"D:\新建文件夹\新建文件夹\新建文件夹\新建文件夹\dsh-mobile\docs\logs\pptx-ref\ref.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst))
